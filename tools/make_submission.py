"""Assemble the Pattern Recognition Letters submission package outside the repository.

    python tools/make_submission.py --out C:/Users/ASUS/Desktop/PRL_submission

Copies the LaTeX source with flat figure paths (the submission system may not keep folders),
the compiled PDF, highlights and graphical abstract, the filled authorship form (local only,
gitignored) and writes the highlights and a draft cover letter as Word files. Rerun after any
change to the manuscript; the output folder is replaced.
"""
import argparse
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / 'paper'
COVER = """Dear Editor,

Please consider our manuscript "{title}" for publication in Pattern Recognition Letters as a research report.

Ground truth for multi-object tracking is increasingly produced by correcting a tracker's output. When the same ground truth is used for evaluation, errors left by the reviewer favour the seed tracker. To our knowledge this seed bias has not been measured in multi-object tracking. Using the public ground truth of MOT17, DanceTrack and FastTrack and a parametric model of the reviewer, we measure the seed's advantage, show that seed-correlated errors reorder trackers more than random errors, and evaluate three inexpensive protocols, including negative results. The analysis plan was pre-registered, and the code and protocol are publicly available.

The manuscript is original, has not been published, and is not under consideration by any other journal. All authors have read and approved it and declare no competing interests. The signed authorship confirmation form is attached.

Sincerely,

Thien Bao Nguyen (corresponding author)
Institute of Intelligent and Interactive Technologies (3ITech), College of Technology and Design
University of Economics Ho Chi Minh City (UEH), Vietnam
baont@ueh.edu.vn
"""


def word(path: Path, title: str, paragraphs=(), bullets=()):
    doc = Document()
    doc.add_heading(title, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    for b in bullets:
        doc.add_paragraph(b, style='List Bullet')
    doc.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', type=Path, required=True)
    args = ap.parse_args()
    out = args.out
    guide = out / '00_HUONG_DAN_NOP_BAI.md'           # hand-written submission guide: kept across reruns
    kept = guide.read_bytes() if guide.exists() else None
    if out.exists():
        shutil.rmtree(out)
    src = out / '01_latex_source'
    src.mkdir(parents=True)
    if kept:
        guide.write_bytes(kept)

    tex = (PAPER / 'main.tex').read_text(encoding='utf-8')
    (src / 'main.tex').write_text(tex.replace('{figures/', '{'), encoding='utf-8')
    for f in ['references.bib', 'prletters.sty', 'elsevier-logo.pdf', *(f'figures/{p.name}' for p in (PAPER / 'figures').glob('*.pdf'))]:
        if not (PAPER / f).exists():
            raise SystemExit(f'missing paper/{f}')
        shutil.copy(PAPER / f, src / Path(f).name)
    shutil.copy(PAPER / 'main.pdf', out / '02_manuscript.pdf')

    items = re.findall(r'\\item (.+)', (PAPER / 'highlights.tex').read_text(encoding='utf-8'))
    items = [i.replace('--', '\u2013').strip() for i in items]
    assert 3 <= len(items) <= 5 and all(len(i) <= 85 for i in items), items
    word(out / '03_highlights.docx', 'Highlights', bullets=items)
    shutil.copy(PAPER / 'graphical_abstract.pdf', out / '04_graphical_abstract.pdf')

    form = PAPER / 'authorship_confirmation_filled.pdf'
    if form.exists():
        shutil.copy(form, out / '05_authorship_confirmation_TO_SIGN.pdf')
    title = re.search(r'\\title\{(.+?)\}\n', tex, re.S).group(1).replace('\\\\', ' ').replace('\n', ' ')
    title = re.sub(r'\s+', ' ', title).strip()
    word(out / '06_cover_letter_DRAFT.docx', 'Cover letter', paragraphs=COVER.format(title=title).split('\n\n'))

    # Overleaf project: main.tex at the root keeps the figures/ folder; compiles with pdfLaTeX + BibTeX.
    with zipfile.ZipFile(out / '07_overleaf_project.zip', 'w', zipfile.ZIP_DEFLATED) as z:
        for f in ['main.tex', 'references.bib', 'prletters.sty', 'elsevier-logo.pdf', 'highlights.tex',
                  'graphical_abstract.pdf', *(f'figures/{p.name}' for p in sorted((PAPER / 'figures').glob('*.pdf')))]:
            z.write(PAPER / f, f)
    print('\n'.join(sorted(str(p.relative_to(out)) for p in out.rglob('*') if p.is_file())))


if __name__ == '__main__':
    main()
